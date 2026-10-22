import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

APP_DIR = Path(__file__).resolve().parents[1] / "compras_app"
sys.path.insert(0, str(APP_DIR))

from config import TEMPLATE_OS
from gerar_os import (
    _celulas_unicas_linha,
    _incluir_ok_nok_por_atividade,
    _inserir_diario_bordo,
    _isolar_layout_na_ultima_pagina,
    _preencher_tabela_processo,
    gerar_os_docx,
)
from os_template import encontrar_linha_cabecalho, mapear_tabelas_os


class ChecklistDiarioTest(unittest.TestCase):
    def test_todos_processos_tem_duas_marcacoes_por_atividade(self):
        doc = Document(TEMPLATE_OS)
        refs = mapear_tabelas_os(doc)
        for nome, idx in refs["processos"].items():
            with self.subTest(processo=nome):
                tabela = doc.tables[idx]
                titulo = tabela.rows[0].cells[0].text
                _preencher_tabela_processo(tabela, [
                    {"atividade": "Conferir montagem", "responsavel": "OPERADOR"},
                    {"atividade": "Testar funcionamento"},
                ])
                cab = encontrar_linha_cabecalho(tabela, "ATIVIDADE")
                self.assertEqual(["#", "ATIVIDADE", "OK", "NOK"],
                                 [c.text for c in _celulas_unicas_linha(tabela.rows[cab])])
                self.assertEqual(["RESPONSAVEL", "DATA", "HORA INÍCIO", "HORA FIM"],
                                 [c.text for c in _celulas_unicas_linha(tabela.rows[1])])
                self.assertEqual(titulo, tabela.rows[0].cells[0].text)
                for n, row in enumerate(tabela.rows[cab + 1:], 1):
                    cells = _celulas_unicas_linha(row)
                    self.assertEqual(4, len(cells))
                    self.assertEqual(str(n), cells[0].text)
                    self.assertEqual(["☐", "☐"], [c.text for c in cells[-2:]])
                    self.assertIsNotNone(row._tr.trPr.find(qn("w:cantSplit")))
                self.assertIn("RESPONSÁVEL: OPERADOR", tabela.rows[cab + 1].cells[1].text)
                xml = tabela._tbl.xml
                _incluir_ok_nok_por_atividade(tabela)
                self.assertEqual(xml, tabela._tbl.xml)

    def test_diario_depois_do_layout_idempotente_e_sem_alterar_tabelas_anteriores(self):
        doc = Document(TEMPLATE_OS)
        layout = doc.tables[mapear_tabelas_os(doc)["layout"]]
        _isolar_layout_na_ultima_pagina(doc, layout)
        antes = [t._tbl.xml for t in doc.tables]
        _inserir_diario_bordo(doc, "3211", {"chassis": "VJ687280"})
        self.assertEqual(antes, [t._tbl.xml for t in doc.tables[:-1]])
        diario = doc.tables[-1]
        self.assertEqual("DIÁRIO DE BORDO DA O.S.", diario.cell(0, 0).text)
        self.assertIn("3211", diario.cell(1, 0).text)
        self.assertIn("VJ687280", diario.cell(1, 0).text)
        self.assertEqual(19, len(diario.rows))
        self.assertEqual(3, len(diario.columns))
        self.assertEqual(16, len(diario.rows[3:]))
        self.assertTrue(diario._tbl.getprevious().xpath(".//w:pageBreakBefore"))
        self.assertLess(doc._body._element.index(layout._tbl), doc._body._element.index(diario._tbl))
        count = len(doc.tables)
        _inserir_diario_bordo(doc, "3211", {"chassis": "VJ687280"})
        self.assertEqual(count, len(doc.tables))

    def test_somente_os_completa_recebe_diario_sem_mudar_outros_documentos(self):
        for modo in ("completa", "producao", "preparacao", "expedicao", "faturamento_direto", "resumida", "mascara"):
            with self.subTest(modo=modo):
                with patch("gerar_os._salvar_documento_os", side_effect=lambda doc, *a, **k: doc):
                    doc = gerar_os_docx("3211", {"chassis": "VJ687280"}, [], {}, {},
                                        composicao_resolvida=[], modo=modo)
                diario = [t for t in doc.tables if "DIÁRIO DE BORDO DA O.S." in t.rows[0].cells[0].text]
                self.assertEqual(1 if modo == "completa" else 0, len(diario))
                self.assertIn("VJ687280", " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells))


if __name__ == "__main__":
    unittest.main()
