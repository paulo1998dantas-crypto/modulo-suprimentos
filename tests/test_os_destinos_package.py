"""Emissão isolada, sem consultar ou escrever na base operacional."""
import copy
import io
import json
import os
import sys
import zipfile
import pytest
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'compras_app'))
os.environ['SUPRIMENTOS_FILE_LOG'] = '0'
import app as application
import supabase_catalog


def test_pacote_nulo_so_no_relatorio_e_pp_todo_na_preparacao(tmp_path):
    catalog = {
        'TRANSF': {'descricao':'TRANSFORMAÇÃO TESTE', 'unidade':'un'},
        'PP': {'grupo':'PP', 'descricao':'PROCESSADO TESTE', 'unidade':'pc'},
        'SUB': {'grupo':'30 - CONJUNTO / KIT', 'descricao':'CJ INTERMEDIÁRIO', 'unidade':'cj'},
        'INSUMO': {'descricao':'MATERIAL DE TESTE', 'unidade':'pc'},
        '30240096': {'descricao':'CJ PRESENTE', 'unidade':'cj'},
        '10240301': {'descricao':'NULO A', 'unidade':'pc'},
        '10240299': {'descricao':'NULO B', 'unidade':'pc'},
        '10320035': {'descricao':'NULO C', 'unidade':'m'},
    }
    bom = {
        'TRANSF':[{'codigo':'PP','quantidade':2}, {'codigo':'30240096','quantidade':1}],
        'PP':[{'codigo':'SUB','quantidade':3}],
        'SUB':[{'codigo':'INSUMO','quantidade':0.8}],
        '30240096':[{'codigo':'10240301','quantidade':2}, {'codigo':'10240299','quantidade':4},
                    {'codigo':'10320035','quantidade':0.5}],
    }
    original = copy.deepcopy(bom)
    material_dir = tmp_path / 'materials'
    material_dir.mkdir()
    values = [
        (application, 'login_enabled', False), (application, 'can', True),
        (application, 'shared_rbac_enabled', False),
        (application.supabase_data, 'enabled', False), (supabase_catalog, 'enabled', False),
        (application, 'atualizar_skus_automatico', {}),
        (application, 'carregar_os_fornecedores', {}),
        (application, 'carregar_os_produtos', catalog), (application, 'carregar_produtos', catalog),
        (application, 'carregar_os_componentes', bom),
        (application, 'carregar_regras_popup_item', []), (application, 'carregar_os_processos', {}),
        (application, 'carregar_relacoes_processo_item', {}), (application, 'get_bom_dir', ''),
        (application, 'limpar_importacao', None), (application, '_preparar_transformacao_documento_os', None),
        (application, 'pasta_os', str(material_dir)),
    ]
    with ExitStack() as stack:
        for target, name, value in values:
            stack.enter_context(patch.object(target, name, return_value=value))
        register = stack.enter_context(patch.object(application, 'registrar_historico', return_value={'id':'TESTE'}))
        application.app.config.update(TESTING=True)
        response = application.app.test_client().post('/gerar_os', data={
            'acao':'imprimir', 'os_numero':'TESTE', 'os_codigo[]':'TRANSF', 'os_qtd[]':'1'})
    assert response.status_code == 200, response.get_data(as_text=True) if response.status_code != 200 else ''
    composition = register.call_args.kwargs['composicao']
    assert bom == original
    assert {line['codigo'] for line in composition} == {'PP','SUB','INSUMO','30240096','10240301','10240299','10320035'}
    assert all(line['setor'] == 'NULO' for line in composition if line['codigo'] in {'10240301','10240299','10320035'})
    assert all(line['setor'] == 'PREPARACAO' for line in composition if line['codigo'] in {'PP','SUB','INSUMO'})
    assert next(line['qtd'] for line in composition if line['codigo'] == 'INSUMO') == pytest.approx(4.8)
    with zipfile.ZipFile(io.BytesIO(response.data)) as package:
        docs = {}
        for name in package.namelist():
            if name.endswith('.docx'):
                doc = Document(io.BytesIO(package.read(name)))
                docs[name] = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        assert len(docs) == 3
        for text in docs.values():
            for sku in ('10240301','10240299','10320035'):
                assert sku not in text
        preparation = next(text for name, text in docs.items() if 'PREPARA' in name.upper())
        shipping = next(text for name, text in docs.items() if 'EXPEDI' in name.upper())
        assert 'PROCESSADO TESTE' in preparation
        assert 'MATERIAL DE TESTE' in preparation
        assert 'MATERIAL DE TESTE' not in shipping
        workbook_name = next(name for name in package.namelist() if name.endswith('.xlsx'))
        workbook = load_workbook(io.BytesIO(package.read(workbook_name)), read_only=True)
        rows = list(workbook['Requisicao'].values)
        headers = rows[1]
        data = [dict(zip(headers, row)) for row in rows[2:]]
        for sku, qtd in [('10240301',2), ('10240299',4), ('10320035',0.5)]:
            row = next(row for row in data if row['CODIGO'] == sku)
            assert row['SETOR'] == 'NULO'
            assert float(str(row['QTD']).replace(',','.')) == qtd
        workbook.close()


def test_pacote_cj_banco_preparacao_insumos_e_capas_expedicao(tmp_path):
    catalog = {
        'TRANSF': {'descricao': 'TRANSFORMACAO TESTE', 'unidade': 'un'},
        '30200036': {'descricao': 'CJ BANCO TESTE', 'categoria': '20 - BANCOS',
                     'grupo': '30 - CONJUNTO / KIT', 'unidade': 'cj'},
        '10200083': {'descricao': 'BCO TESTE C/TRILHO', 'categoria': '20 - BANCOS',
                     'grupo': '10 - INSUMO', 'unidade': 'pc'},
        '10220001': {'descricao': 'CAPA TESTE', 'categoria': '22 - PECAS BCO',
                     'grupo': '10 - INSUMO', 'unidade': 'pc'},
    }
    bom = {
        'TRANSF': [{'codigo': '30200036', 'quantidade': 1}],
        '30200036': [{'codigo': '10200083', 'quantidade': 3}],
        '10200083': [{'codigo': '10220001', 'quantidade': 2}],
    }
    original = copy.deepcopy(bom)
    material_dir = tmp_path / 'materials'
    material_dir.mkdir()
    values = [
        (application, 'login_enabled', False), (application, 'can', True),
        (application, 'shared_rbac_enabled', False),
        (application.supabase_data, 'enabled', False), (supabase_catalog, 'enabled', False),
        (application, 'atualizar_skus_automatico', {}),
        (application, 'carregar_os_fornecedores', {}),
        (application, 'carregar_os_produtos', catalog), (application, 'carregar_produtos', catalog),
        (application, 'carregar_os_componentes', bom),
        (application, 'carregar_regras_popup_item', []), (application, 'carregar_os_processos', {}),
        (application, 'carregar_relacoes_processo_item', {}), (application, 'get_bom_dir', ''),
        (application, 'limpar_importacao', None), (application, '_preparar_transformacao_documento_os', None),
        (application, 'pasta_os', str(material_dir)),
    ]
    with ExitStack() as stack:
        for target, name, value in values:
            stack.enter_context(patch.object(target, name, return_value=value))
        register = stack.enter_context(patch.object(application, 'registrar_historico', return_value={'id': 'TESTE'}))
        application.app.config.update(TESTING=True)
        response = application.app.test_client().post('/gerar_os', data={
            'acao': 'imprimir', 'os_numero': 'TESTE', 'os_codigo[]': 'TRANSF', 'os_qtd[]': '1'})
    assert response.status_code == 200
    assert bom == original
    expected = {'30200036': (1, 'PREPARACAO'), '10200083': (3, 'EXPEDICAO'),
                '10220001': (6, 'EXPEDICAO')}
    composition = {line['codigo']: line for line in register.call_args.kwargs['composicao']}
    assert set(composition) == set(expected)
    for sku, (qtd, setor) in expected.items():
        assert composition[sku]['qtd'] == qtd
        assert composition[sku]['setor'] == setor
    with zipfile.ZipFile(io.BytesIO(response.data)) as package:
        docs = {}
        for name in package.namelist():
            if name.endswith('.docx'):
                doc = Document(io.BytesIO(package.read(name)))
                docs[name] = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        preparation = next(text for name, text in docs.items() if 'PREPARA' in name.upper())
        shipping = next(text for name, text in docs.items() if 'EXPEDI' in name.upper())
        assert 'CJ BANCO TESTE' in preparation and 'CJ BANCO TESTE' not in shipping
        for sku in ('10200083', '10220001'):
            assert sku in shipping and sku not in preparation
        workbook_name = next(name for name in package.namelist() if name.endswith('.xlsx'))
        workbook = load_workbook(io.BytesIO(package.read(workbook_name)), read_only=True)
        rows = list(workbook['Requisicao'].values)
        data = [dict(zip(rows[1], row)) for row in rows[2:]]
        data = {row['CODIGO']: row for row in data}
        for sku, (qtd, setor) in expected.items():
            assert data[sku]['SETOR'] == setor
            assert float(str(data[sku]['QTD']).replace(',', '.')) == qtd
        workbook.close()


def test_docx_nao_imprime_nulo_manual_nem_recalcula_composicao():
    from gerar_os import gerar_os_docx
    rows = [{'codigo':'SKU-OCULTO', 'descricao':'NÃO IMPRIMIR', 'qtd':7, 'setor':'NULO', 'setor_manual':True}]
    original = copy.deepcopy(rows)
    for mode in ('completa', 'preparacao', 'expedicao', 'faturamento_direto', 'resumida', 'mascara'):
        with patch('gerar_os._salvar_documento_os', side_effect=lambda doc, *a, **k: doc):
            doc = gerar_os_docx('TESTE', {}, rows, {}, {}, composicao_resolvida=rows, modo=mode)
        text = ' '.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        assert 'SKU-OCULTO' not in text
        assert 'NÃO IMPRIMIR' not in text
        assert rows == original


def test_os_completa_mantem_composicao_dos_skus_nulo():
    from gerar_os import gerar_os_docx

    item_pai = {
        'codigo': '30240032', 'descricao': 'ACESSORIO CJ EXTINTOR COM SUPORTE ABC 4KG',
        'qtd': 1, 'unidade': 'cj', 'setor': 'NULO',
    }
    componentes = {
        '30240032': [
            {'codigo': '10240044', 'descricao': 'ACESSORIO EXTINTOR ABC 4KG'},
            {'codigo': '10240045', 'descricao': 'ACESSORIO SUPORTE EXTINTOR ABC 4KG'},
        ],
    }
    composicao = [
        {'item': '30240032', 'codigo': '10240044', 'descricao': 'ACESSORIO EXTINTOR ABC 4KG',
         'qtd': 1, 'unidade': 'pc', 'level': 0, 'setor': 'EXPEDICAO'},
        {'item': '30240032', 'codigo': '10240045', 'descricao': 'ACESSORIO SUPORTE EXTINTOR ABC 4KG',
         'qtd': 1, 'unidade': 'pc', 'level': 0, 'setor': 'EXPEDICAO'},
    ]

    with patch('gerar_os._salvar_documento_os', side_effect=lambda doc, *a, **k: doc):
        doc = gerar_os_docx('TESTE', {}, [item_pai], componentes, {},
                            composicao_resolvida=composicao, modo='completa')

    text = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert '30240032' not in text
    assert '10240044' in text and 'ACESSORIO EXTINTOR ABC 4KG' in text
    assert '10240045' in text and 'ACESSORIO SUPORTE EXTINTOR ABC 4KG' in text
