import os
import sys
import unittest
from pathlib import Path


APP_DIR = Path(__file__).resolve().parents[1] / "compras_app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))
os.environ["SUPRIMENTOS_FILE_LOG"] = "0"

import app as app_module  # noqa: E402
import supabase_data  # noqa: E402
from docx import Document  # noqa: E402
from gerar_os import _inserir_dados_cliente_faturamento  # noqa: E402


class WorkOrderCustomerBillingTests(unittest.TestCase):
    def test_direct_billing_refreshes_only_customer_contacts_by_registry_id(self):
        dados = {
            "cliente": "CLIENTE EXEMPLO",
            "cliente_cadastro": {
                "cnpj_cpf": "12.345.678/0001-90",
                "razao_social": "CLIENTE EXEMPLO LTDA",
                "telefone": "",
                "email": "antigo@example.com",
            },
        }
        clientes = {
            "outro": {
                "cliente": "CLIENTE EXEMPLO",
                "telefone": "0000000000",
                "email": "outro@example.com",
            },
            "12.345.678/0001-90": {
                "cliente": "CLIENTE EXEMPLO",
                "telefone": "1133334444",
                "celular": "11999998888",
                "whatsapp": "11999998888",
                "email": "contato@example.com",
                "site": "exemplo.com.br",
            },
        }

        cadastro = app_module._contatos_cliente_faturamento_direto(dados, clientes)

        self.assertEqual("CLIENTE EXEMPLO LTDA", cadastro["razao_social"])
        self.assertEqual("1133334444", cadastro["telefone"])
        self.assertEqual("11999998888", cadastro["celular"])
        self.assertEqual("11999998888", cadastro["whatsapp"])
        self.assertEqual("contato@example.com", cadastro["email"])
        self.assertEqual("exemplo.com.br", cadastro["site"])
        self.assertEqual("antigo@example.com", dados["cliente_cadastro"]["email"])

    def test_customer_resolves_by_registry_key_and_keeps_full_snapshot(self):
        customers = {
            "12.345.678/0001-90": {
                "cliente": "CLIENTE EXEMPLO",
                "nome_fantasia": "CLIENTE EXEMPLO",
                "razao_social": "CLIENTE EXEMPLO LTDA",
                "cnpj_cpf": "12.345.678/0001-90",
                "ie": "123456",
                "logradouro": "Rua Um",
                "logradouro_numero": "25",
                "complemento": "Galpão 2",
                "bairro": "Centro",
                "cidade": "Campinas",
                "uf": "SP",
                "cep": "13000-000",
                "email": "fiscal@example.com",
                "limite_credito": 1000,
                "payload": {"campo_extra": "valor"},
            }
        }

        cadastro = app_module._resolver_cadastro_cliente_os("CLIENTE EXEMPLO", customers)
        snapshot = app_module._snapshot_cadastro_cliente_os(cadastro)

        self.assertEqual("12.345.678/0001-90", snapshot["cnpj_cpf"])
        self.assertEqual("123456", snapshot["ie"])
        self.assertEqual("Galpão 2", snapshot["complemento"])
        self.assertEqual(1000, snapshot["limite_credito"])
        self.assertEqual({"campo_extra": "valor"}, snapshot["payload"])

    def test_supabase_customer_adapter_keeps_billing_fields(self):
        cliente = supabase_data._pessoa_to_legacy(
            {
                "identificador": "12345678900",
                "nome_fantasia": "CLIENTE",
                "razao_social": "CLIENTE LTDA",
                "cnpj_cpf": "12345678900",
                "ie": "ISENTO",
                "logradouro": "Avenida Central",
                "logradouro_numero": "100",
                "complemento": "Sala 4",
                "celular": "11999990000",
                "email": "fiscal@example.com",
                "cliente": True,
            },
            "cliente",
        )

        self.assertEqual("CLIENTE", cliente["cliente"])
        self.assertEqual("CLIENTE LTDA", cliente["razao_social"])
        self.assertEqual("ISENTO", cliente["ie"])
        self.assertEqual("100", cliente["logradouro_numero"])
        self.assertEqual("Sala 4", cliente["complemento"])
        self.assertEqual("11999990000", cliente["celular"])

    def test_direct_billing_document_receives_registered_customer_details(self):
        doc = Document()
        _inserir_dados_cliente_faturamento(
            doc,
            {
                "cliente": "CLIENTE EXEMPLO",
                "cliente_cadastro": {
                    "nome_fantasia": "CLIENTE EXEMPLO",
                    "razao_social": "CLIENTE EXEMPLO LTDA",
                    "cnpj_cpf": "12.345.678/0001-90",
                    "ie": "123456",
                    "logradouro": "Rua Um",
                    "logradouro_numero": "25",
                    "complemento": "Galpão 2",
                    "bairro": "Centro",
                    "cidade": "Campinas",
                    "uf": "SP",
                    "cep": "13000-000",
                    "telefone": "1133334444",
                    "celular": "11999998888",
                    "whatsapp": "11999997777",
                    "email": "fiscal@example.com",
                    "site": "exemplo.com.br",
                },
            },
            {},
        )
        texto = " ".join(
            [paragraph.text for paragraph in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )

        self.assertIn("DADOS CADASTRAIS DO CLIENTE PARA FATURAMENTO DIRETO", texto)
        self.assertIn("12.345.678/0001-90", texto)
        self.assertIn("Rua Um, 25, Galpão 2", texto)
        self.assertIn("1133334444", texto)
        self.assertIn("11999998888", texto)
        self.assertIn("11999997777", texto)
        self.assertIn("fiscal@example.com", texto)
        self.assertIn("exemplo.com.br", texto)
        contatos = {
            row.cells[0].text: row.cells[1].text
            for table in doc.tables for row in table.rows
        }
        self.assertEqual("1133334444", contatos["Telefone"])
        self.assertEqual("11999998888", contatos["Celular"])
        self.assertEqual("11999997777", contatos["WhatsApp"])
        self.assertEqual("fiscal@example.com", contatos["E-mail"])


if __name__ == "__main__":
    unittest.main()
