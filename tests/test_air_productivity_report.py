import io
import os
import sys
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook


APP_DIR = Path(__file__).resolve().parents[1] / "compras_app"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))
os.environ["SUPRIMENTOS_FILE_LOG"] = "0"

import app as app_module  # noqa: E402
from air_productivity_report import (  # noqa: E402
    build_air_productivity_workbook,
    prepare_air_productivity_data,
)


def order(os_number, item, supplier, finished, chassis, **extra):
    return {
        "numero_os": os_number,
        "item_number": item,
        "ar_condicionado": supplier,
        "termino_producao": finished,
        "chassi": chassis,
        "situacao": "FINALIZADA",
        "modelo_veicular": "SPRINTER 417",
        "transformacao": "TRANSFORMAÇÃO DE TESTE",
        **extra,
    }


class AirProductivityReportTests(unittest.TestCase):
    def setUp(self):
        app_module.app.config.update(TESTING=True)
        self.client = app_module.app.test_client()

    def test_uses_current_completion_supplier_and_counts_duplicate_chassis_once(self):
        report = prepare_air_productivity_data([
            order("100", 3000, "CLIM", "2026-09-01T15:00:00-03:00", "CHASSI-A"),
            order("101", 3001, "GE", "2026-09-03", "CHASSI-B"),
            order("102", 3002, "GE", "2026-09-04", "CHASSI-A"),
            order("103", 3003, "CLIM", "2026-09-05", "CHASSI-C", situacao="EM PRODUÇÃO"),
            order("104", 3004, "OUTRO", "2026-09-07", "CHASSI-D"),
            order("105", 3005, "CLIM", "", "CHASSI-E"),
        ], target_per_supplier=2)

        self.assertEqual(2, report["unique_count"])
        self.assertEqual(1, report["duplicate_count"])
        self.assertEqual(date(2026, 9, 1), report["period_start"])
        self.assertEqual(date(2026, 9, 3), report["period_end"])
        self.assertEqual(1, report["monthly"][0]["clim"])
        self.assertEqual(1, report["monthly"][0]["ge"])
        duplicate = next(row for row in report["rows"] if row["numero_os"] == "102")
        self.assertFalse(duplicate["conta"])

    def test_workbook_has_three_views_formulas_and_the_new_target(self):
        report = prepare_air_productivity_data([
            order("100", 3000, "CLIM", "2026-09-01", "CHASSI-A"),
            order("101", 3001, "GE", "2026-09-03", "CHASSI-B"),
        ], target_per_supplier=2)
        workbook = build_air_productivity_workbook(report)
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        exported = load_workbook(output, data_only=False)
        self.assertEqual(["Resumo Mensal", "Produção Diária", "Concluídos"], exported.sheetnames)
        self.assertEqual(2, exported["Resumo Mensal"]["I5"].value)
        self.assertTrue(exported["Resumo Mensal"]["C10"].value.startswith("=SUMIFS"))
        self.assertTrue(exported["Produção Diária"]["D5"].value.startswith("=COUNTIFS"))
        self.assertEqual("dd/mm/yyyy", exported["Concluídos"]["D5"].number_format)
        self.assertTrue(exported.calculation.fullCalcOnLoad)
        exported.close()

    def test_completion_timestamp_is_grouped_by_brazilian_operational_date(self):
        report = prepare_air_productivity_data([
            order("100", 3000, "CLIM", "2026-10-07T02:00:00Z", "CHASSI-A"),
        ])
        self.assertEqual(date(2026, 10, 6), report["period_start"])

    def test_pcp_export_reads_fresh_mes_data_and_is_linked_from_management_screen(self):
        source = [order("100", 3000, "CLIM", "2026-09-01", "CHASSI-A")]
        with (
            patch.object(app_module, "login_enabled", return_value=False),
            patch.object(app_module, "erp_feature_enabled", return_value=True),
            patch.object(app_module, "can", return_value=True),
            patch.object(app_module, "_erp_mes_all_work_orders", return_value=source) as loader,
        ):
            first = self.client.get("/erp/relatorios/produtividade-fornecedor-ar.xlsx")
            second = self.client.get("/erp/relatorios/produtividade-fornecedor-ar.xlsx")

        self.assertEqual(200, first.status_code)
        self.assertEqual(200, second.status_code)
        self.assertEqual(2, loader.call_count)
        self.assertIn("Produtividade_Ar_", first.headers["Content-Disposition"])
        template = (APP_DIR / "templates" / "erp_gestao_os.html").read_text(encoding="utf-8")
        self.assertIn("/erp/relatorios/produtividade-fornecedor-ar.xlsx", template)

    def test_mes_work_order_loader_paginates_without_reusing_a_cache(self):
        page = [{"numero_os": str(index)} for index in range(1500)]
        with patch.object(
            app_module,
            "_erp_mes_request",
            side_effect=[{"orders": page}, {"orders": [{"numero_os": "1500"}]}],
        ) as request_mes:
            orders = app_module._erp_mes_all_work_orders()

        self.assertEqual(1501, len(orders))
        self.assertEqual(
            ["work-orders?limit=1500&offset=0", "work-orders?limit=1500&offset=1500"],
            [call.args[0] for call in request_mes.call_args_list],
        )


if __name__ == "__main__":
    unittest.main()
